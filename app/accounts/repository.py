from __future__ import annotations

import csv
import json
import threading
from dataclasses import asdict, dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, Iterable, Mapping


class AccountRepositoryError(RuntimeError):
    """Raised when the account source cannot be read or written."""


class AccountStatus(str, Enum):
    READY = "READY"
    IN_USE = "IN_USE"
    DONE = "DONE"
    FAILED = "FAILED"
    DISABLED = "DISABLED"

    @classmethod
    def normalize(cls, value: Any, default: "AccountStatus" = READY) -> "AccountStatus":
        if isinstance(value, cls):
            return value
        text = str(value or "").strip().upper()
        aliases = {
            "": default,
            "NEW": cls.READY,
            "PENDING": cls.READY,
            "AVAILABLE": cls.READY,
            "RUNNING": cls.IN_USE,
            "BUSY": cls.IN_USE,
            "SUCCESS": cls.DONE,
            "COMPLETED": cls.DONE,
            "ERROR": cls.FAILED,
            "FAIL": cls.FAILED,
            "OFF": cls.DISABLED,
            "BLOCKED": cls.DISABLED,
        }
        if text in aliases:
            return aliases[text]
        try:
            return cls(text)
        except ValueError:
            return default


@dataclass(slots=True)
class Account:
    id: str
    username: str
    password: str = ""
    status: AccountStatus = AccountStatus.READY
    attempts: int = 0
    assigned_worker: str | None = None
    last_error: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_variables(self) -> dict[str, Any]:
        """Safe scenario variable mapping. Password is included for local automation use."""
        values: dict[str, Any] = {
            "id": self.id,
            "username": self.username,
            "password": self.password,
            "status": self.status.value,
            "attempts": self.attempts,
        }
        values.update(self.metadata)
        return values

    def to_record(self) -> dict[str, Any]:
        record = {
            "id": self.id,
            "username": self.username,
            "password": self.password,
            "status": self.status.value,
            "attempts": self.attempts,
            "assigned_worker": self.assigned_worker or "",
            "last_error": self.last_error or "",
        }
        for key, value in self.metadata.items():
            if key not in record:
                record[key] = value
        return record


_USERNAME_ALIASES = ("username", "user", "account", "email", "login", "tai_khoan", "tài_khoản")
_PASSWORD_ALIASES = ("password", "pass", "pwd", "mat_khau", "mật_khẩu")
_ID_ALIASES = ("id", "account_id", "uid", "stt")
_STATUS_ALIASES = ("status", "state", "trang_thai", "trạng_thái")
_ATTEMPTS_ALIASES = ("attempts", "retry", "retries", "so_lan_thu", "số_lần_thử")
_RESERVED = {
    *_USERNAME_ALIASES,
    *_PASSWORD_ALIASES,
    *_ID_ALIASES,
    *_STATUS_ALIASES,
    *_ATTEMPTS_ALIASES,
    "assigned_worker",
    "last_error",
}


def _normalize_key(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _first_value(record: Mapping[str, Any], aliases: Iterable[str], default: Any = None) -> Any:
    for alias in aliases:
        if alias in record and record[alias] not in (None, ""):
            return record[alias]
    return default


class AccountRepository:
    """Read/write account lists from XLSX, CSV or JSON files."""

    def __init__(self, source: str | Path) -> None:
        self.path = Path(source).expanduser()
        self._lock = threading.RLock()

    @property
    def exists(self) -> bool:
        return self.path.is_file()

    def load(self) -> list[Account]:
        with self._lock:
            if not self.path.is_file():
                raise FileNotFoundError(f"Account source not found: {self.path}")

            suffix = self.path.suffix.lower()
            if suffix in {".xlsx", ".xlsm"}:
                rows = self._load_xlsx()
            elif suffix == ".csv":
                rows = self._load_csv()
            elif suffix == ".json":
                rows = self._load_json()
            else:
                raise AccountRepositoryError(
                    f"Unsupported account file format: {self.path.suffix}. Use .xlsx, .csv or .json"
                )

            accounts: list[Account] = []
            seen_ids: set[str] = set()
            for index, raw in enumerate(rows, start=1):
                normalized = {_normalize_key(key): value for key, value in raw.items() if key is not None}
                username = str(_first_value(normalized, _USERNAME_ALIASES, "") or "").strip()
                if not username:
                    continue

                account_id = str(_first_value(normalized, _ID_ALIASES, index) or index).strip()
                if account_id in seen_ids:
                    account_id = f"{account_id}_{index}"
                seen_ids.add(account_id)

                password = str(_first_value(normalized, _PASSWORD_ALIASES, "") or "")
                status = AccountStatus.normalize(_first_value(normalized, _STATUS_ALIASES, AccountStatus.READY.value))
                attempts_raw = _first_value(normalized, _ATTEMPTS_ALIASES, 0)
                try:
                    attempts = max(0, int(attempts_raw or 0))
                except (TypeError, ValueError):
                    attempts = 0

                metadata = {
                    key: value
                    for key, value in normalized.items()
                    if key not in _RESERVED and value not in (None, "")
                }
                accounts.append(
                    Account(
                        id=account_id,
                        username=username,
                        password=password,
                        status=status,
                        attempts=attempts,
                        assigned_worker=(str(normalized.get("assigned_worker") or "").strip() or None),
                        last_error=(str(normalized.get("last_error") or "").strip() or None),
                        metadata=metadata,
                    )
                )
            return accounts

    def save(self, accounts: Iterable[Account]) -> Path:
        with self._lock:
            records = [account.to_record() for account in accounts]
            self.path.parent.mkdir(parents=True, exist_ok=True)
            suffix = self.path.suffix.lower()
            if suffix in {".xlsx", ".xlsm"}:
                self._save_xlsx(records)
            elif suffix == ".csv":
                self._save_csv(records)
            elif suffix == ".json":
                self._save_json(records)
            else:
                raise AccountRepositoryError(f"Unsupported account file format: {self.path.suffix}")
            return self.path

    def _load_csv(self) -> list[dict[str, Any]]:
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle)]

    def _load_json(self) -> list[dict[str, Any]]:
        with self.path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        if isinstance(data, dict):
            data = data.get("accounts", [])
        if not isinstance(data, list):
            raise AccountRepositoryError("JSON account source must be a list or {'accounts': [...]} mapping")
        return [dict(row) for row in data if isinstance(row, Mapping)]

    def _load_xlsx(self) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise AccountRepositoryError("openpyxl is required to read .xlsx account files") from exc

        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            try:
                header_row = next(iterator)
            except StopIteration:
                return []
            headers = [_normalize_key(value) for value in header_row]
            rows: list[dict[str, Any]] = []
            for values in iterator:
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]})
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _headers(records: list[dict[str, Any]]) -> list[str]:
        preferred = ["id", "username", "password", "status", "attempts", "assigned_worker", "last_error"]
        headers = list(preferred)
        for record in records:
            for key in record:
                if key not in headers:
                    headers.append(key)
        return headers

    def _save_csv(self, records: list[dict[str, Any]]) -> None:
        headers = self._headers(records)
        with self.path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)

    def _save_json(self, records: list[dict[str, Any]]) -> None:
        with self.path.open("w", encoding="utf-8") as handle:
            json.dump({"accounts": records}, handle, ensure_ascii=False, indent=2)

    def _save_xlsx(self, records: list[dict[str, Any]]) -> None:
        try:
            from openpyxl import Workbook
        except ImportError as exc:  # pragma: no cover
            raise AccountRepositoryError("openpyxl is required to write .xlsx account files") from exc

        headers = self._headers(records)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "accounts"
        sheet.append(headers)
        for record in records:
            sheet.append([record.get(header, "") for header in headers])
        workbook.save(self.path)
